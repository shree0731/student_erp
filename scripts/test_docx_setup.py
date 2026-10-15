import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def test_docx():
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Page Border
    sectPr = section._sectPr
    border_xml = parse_xml(
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        f'<w:top w:val="single" w:sz="8" w:space="24" w:color="000000"/>'
        f'<w:left w:val="single" w:sz="8" w:space="24" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="24" w:color="000000"/>'
        f'<w:right w:val="single" w:sz="8" w:space="24" w:color="000000"/>'
        f'</w:pgBorders>'
    )
    sectPr.append(border_xml)

    # Centered Page Number in footer
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = f_p.add_run()
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    f_run._r.append(fldSimple)

    p = doc.add_paragraph("EASWARI ENGINEERING COLLEGE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save("test_border.docx")
    print("test_border.docx created successfully")

if __name__ == "__main__":
    test_docx()
